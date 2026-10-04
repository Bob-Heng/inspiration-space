"""TASK-019 导出单元测试：条目行格式、原始观点/分类观点 docx 生成与回读。"""

from datetime import date
from io import BytesIO

from docx import Document

from app.domain.docx_export import (
    build_classified_docx,
    build_viewpoints_docx,
    format_viewpoint_line,
)
from app.models import Viewpoint


def _vp(**overrides) -> Viewpoint:
    data = {
        "id": 7,
        "type": "raw",
        "content": "正文首段。",
        "source_inspiration_id": None,
        "source_date": date(2026, 5, 2),
        "layer": "fa",
        "domain": "经济",
        "circle": None,
        "discipline": "经济学",
        "scene": "文旅消费",
        "status": "accepted",
    }
    data.update(overrides)
    return Viewpoint(**data)


def _read_paragraphs(content: bytes) -> list[str]:
    doc = Document(BytesIO(content))
    return [p.text for p in doc.paragraphs]


class TestFormatViewpointLine:
    def test_完整格式(self):
        assert (
            format_viewpoint_line(_vp())
            == "7.[5.2][法][经济//经济学/文旅消费]正文首段。"
        )

    def test_空位留空(self):
        line = format_viewpoint_line(
            _vp(source_date=None, layer=None, domain=None, discipline=None, scene=None)
        )
        assert line == "7.[][][///]正文首段。"

    def test_跨段正文首行只取首段(self):
        line = format_viewpoint_line(_vp(content="首段。\n第二段。\n第三段。"))
        assert line.endswith("首段。")
        assert "第二段" not in line


class TestBuildViewpointsDocx:
    def test_结构与续段还原(self):
        content = build_viewpoints_docx(
            [_vp(content="首段。\n续段。"), _vp(id=8)]
        )
        paragraphs = _read_paragraphs(content)
        assert paragraphs[0] == "原始观点"
        assert paragraphs[2] == "7.[5.2][法][经济//经济学/文旅消费]首段。"
        assert paragraphs[3] == "续段。"
        assert paragraphs[4].startswith("8.[5.2][法][经济//经济学/文旅消费]正文首段。")

    def test_生成的docx可正常打开(self):
        content = build_viewpoints_docx([_vp()])
        doc = Document(BytesIO(content))
        assert len(doc.paragraphs) == 3  # 标题 + 说明 + 1 条


class TestBuildClassifiedDocx:
    def test_三节分组(self):
        groups = {
            "dao": [_vp(id=1, layer="dao")],
            "fa": [],
            "shu": [_vp(id=2, layer="shu")],
        }
        paragraphs = _read_paragraphs(build_classified_docx(groups))
        assert paragraphs[0] == "分类观点"
        dao_at = paragraphs.index("道")
        shu_at = paragraphs.index("术")
        assert dao_at < shu_at
        assert paragraphs[dao_at + 1].startswith("1.[5.2][道]")
        assert paragraphs[shu_at + 1].startswith("2.[5.2][术]")
        # 法节为空：法标题后紧邻术标题
        assert paragraphs[paragraphs.index("法") + 1] == "术"
